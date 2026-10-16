import copy
import os
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.hyperlink import Hyperlink

RAIZ = Path(__file__).resolve().parents[2]  # carpeta cerdas-exel
RUTA_PLANILLA = RAIZ / "exel" / "planilla.xlsx"

# En un servidor con disco persistente, DATA_DIR apunta a ese disco (ej. /var/data)
CARPETA_DATOS = Path(os.getenv("DATA_DIR") or RAIZ / "exel" / "salida")
RUTA_SALIDA = CARPETA_DATOS / "planilla_actualizada.xlsx"

PRIMERA_FILA_DATOS = 5      # en cada hoja de cerda
PRIMERA_FILA_INDICE = 13    # en INDICE_CERDAS
ULTIMA_FILA_INDICE = 1000
NUM_COLUMNAS = 8            # columnas A a H


def _poner_enlace(celda, hoja_destino, texto):
    """Enlace interno real (no fórmula): funciona al abrir el archivo descargado,
    incluso en 'Vista protegida', donde las fórmulas todavía no se calculan."""
    celda.value = texto
    celda.hyperlink = Hyperlink(ref=celda.coordinate, location=f"'{hoja_destino}'!A1", display=texto)


def reparar_enlaces(wb):
    """Deja funcionando todos los enlaces: 'Abrir' del índice y 'VOLVER AL ÍNDICE' de cada hoja.
    Si una cerda figura en el índice pero no tiene hoja, se la crea vacía."""
    wb.calculation.fullCalcOnLoad = True   # Excel recalcula las fórmulas del buscador al abrir
    indice = wb["INDICE_CERDAS"]
    for fila in range(PRIMERA_FILA_INDICE, indice.max_row + 1):
        nro = indice.cell(fila, 1).value
        if nro is None:
            continue
        _obtener_hoja(wb, nro)
        _poner_enlace(indice.cell(fila, 2), str(nro), "Abrir")
    for nombre in wb.sheetnames:
        if nombre != "INDICE_CERDAS":
            _poner_enlace(wb[nombre]["A1"], "INDICE_CERDAS", "← VOLVER AL ÍNDICE")


def _campos(d):
    """Valores en el orden exacto de las columnas A..H."""
    if d.no_gesto:
        return [d.fecha_servicio, d.padrillo, None, None, None, None, None, None]
    return [d.fecha_servicio, d.padrillo, d.fecha_parto_probable,
            d.fecha_parto, d.nacidos_vivos, d.nacidos_totales,
            d.fecha_destete, d.destetados]


def _obtener_hoja(wb, nro):
    nombre = str(nro)
    if nombre in wb.sheetnames:
        return wb[nombre]

    if "PLANTILLA" in wb.sheetnames:
        base = wb["PLANTILLA"]
    else:
        # sin PLANTILLA: usa la primera hoja de cerda como molde
        base = next(wb[n] for n in wb.sheetnames if n != "INDICE_CERDAS")

    hoja = wb.copy_worksheet(base)
    hoja.title = nombre
    hoja["A2"] = f"CERDA N. {nro}"
    hoja["A3"] = f"CERDA N. {nro}"
    _poner_enlace(hoja["A1"], "INDICE_CERDAS", "← VOLVER AL ÍNDICE")
    for fila in range(PRIMERA_FILA_DATOS, hoja.max_row + 1):   # vacía los datos copiados
        for col in range(1, NUM_COLUMNAS + 1):
            hoja.cell(fila, col).value = None
    return hoja


def _fila_tiene_datos(hoja, fila):
    return any(hoja.cell(fila, c).value is not None for c in range(1, NUM_COLUMNAS + 1))


def _fila_vacia(hoja):
    fila = PRIMERA_FILA_DATOS
    while _fila_tiene_datos(hoja, fila):
        fila += 1
    return fila


def _buscar_servicio(hoja, fecha):
    """Fila que ya tiene esa fecha de servicio (para completarla)."""
    if fecha is None:
        return None
    for fila in range(PRIMERA_FILA_DATOS, hoja.max_row + 1):
        v = hoja.cell(fila, 1).value
        if isinstance(v, datetime) and v.date() == fecha.date():
            return fila
    return None


def _copiar_formato(hoja, fila):
    if fila == PRIMERA_FILA_DATOS:
        return
    for c in range(1, NUM_COLUMNAS + 1):
        hoja.cell(fila, c)._style = copy.copy(hoja.cell(PRIMERA_FILA_DATOS, c)._style)
    hoja.row_dimensions[fila].height = hoja.row_dimensions[PRIMERA_FILA_DATOS].height


def _actualizar_indice(wb, nro, hoja):
    indice = wb["INDICE_CERDAS"]
    fila_idx = None
    for fila in range(PRIMERA_FILA_INDICE, ULTIMA_FILA_INDICE + 1):
        v = indice.cell(fila, 1).value
        if v == nro:
            fila_idx = fila
            break
        if v is None:
            fila_idx = fila
            indice.cell(fila, 1).value = nro
            break
    if fila_idx is None:
        raise ValueError("El índice está lleno (A13:A1000)")

    _poner_enlace(indice.cell(fila_idx, 2), str(nro), "Abrir")

    registros, ultimo = 0, None
    for fila in range(PRIMERA_FILA_DATOS, hoja.max_row + 1):
        if _fila_tiene_datos(hoja, fila):
            registros += 1
        v = hoja.cell(fila, 1).value
        if isinstance(v, datetime) and (ultimo is None or v > ultimo):
            ultimo = v
    indice.cell(fila_idx, 3).value = registros
    indice.cell(fila_idx, 4).value = ultimo
    indice.cell(fila_idx, 4).number_format = "dd/mm/yyyy"


def guardar_registros(lista_datos):
    """Recibe una lista de DatosCerdo y escribe todo en el Excel."""
    # Si ya existe una salida, se sigue sobre ella (así no se pierde lo guardado antes).
    # El archivo original (exel/planilla.xlsx) nunca se modifica.
    base = RUTA_SALIDA if RUTA_SALIDA.exists() else RUTA_PLANILLA
    wb = load_workbook(base)
    for d in lista_datos:
        hoja = _obtener_hoja(wb, d.cerda)
        fila = _buscar_servicio(hoja, d.fecha_servicio) or _fila_vacia(hoja)
        _copiar_formato(hoja, fila)
        for col, valor in enumerate(_campos(d), start=1):
            if valor is not None:  # no pisa datos existentes
                hoja.cell(fila, col).value = valor
        _actualizar_indice(wb, d.cerda, hoja)
    reparar_enlaces(wb)
    RUTA_SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(RUTA_SALIDA)
    return RUTA_SALIDA


# ----------------------------------------------------------------------
#  Excel de trabajo: cargar, consultar, reiniciar y buscar
# ----------------------------------------------------------------------
def archivo_activo():
    """El Excel con el que se trabaja: la copia cargada/actualizada, o el original."""
    return RUTA_SALIDA if RUTA_SALIDA.exists() else RUTA_PLANILLA


def _texto_fecha(v):
    return v.strftime("%d/%m/%Y") if isinstance(v, datetime) else v


def _validar_formato(wb):
    """Devuelve un mensaje de error, o None si el Excel sirve."""
    if "INDICE_CERDAS" not in wb.sheetnames:
        return "El Excel no tiene la hoja INDICE_CERDAS."
    hojas = [n for n in wb.sheetnames if n not in ("INDICE_CERDAS", "PLANTILLA")]
    if not hojas:
        return "El Excel no tiene hojas de cerdas."
    for n in hojas[:3]:
        fila = next(wb[n].iter_rows(min_row=4, max_row=4, values_only=True), ())
        a = str(fila[0] or "").lower() if len(fila) > 0 else ""
        b = str(fila[1] or "").lower() if len(fila) > 1 else ""
        if "servicio" not in a or "padrillo" not in b:
            return (f"La hoja {n} no tiene el orden de columnas esperado "
                    "(Fecha de servicio, Padrillo, Fecha parición probable, Fecha parto, "
                    "Lech. nacidos vivos, Lech. nacidos totales, Fecha destete, Lech. dest.). "
                    "Usá el Excel en el formato nuevo.")
    return None


def cargar_excel(contenido):
    """Guarda un Excel subido como la copia de trabajo. Lanza ValueError si no sirve."""
    RUTA_SALIDA.parent.mkdir(parents=True, exist_ok=True)
    tmp = RUTA_SALIDA.with_name("_subiendo.xlsx")
    tmp.write_bytes(contenido)
    try:
        wb = load_workbook(tmp, read_only=True)
        error = _validar_formato(wb)
        hojas = len([n for n in wb.sheetnames if n not in ("INDICE_CERDAS", "PLANTILLA")])
        wb.close()
    except ValueError:
        raise
    except Exception:
        tmp.unlink(missing_ok=True)
        raise ValueError("No se pudo abrir el archivo. ¿Es un Excel (.xlsx)?")
    if error:
        tmp.unlink(missing_ok=True)
        raise ValueError(error)
    # Se vuelven a escribir los enlaces como enlaces reales (los del original eran fórmulas)
    wb = load_workbook(tmp)
    reparar_enlaces(wb)
    wb.save(tmp)
    os.replace(tmp, RUTA_SALIDA)
    return {"hojas": hojas}


def estado():
    activo = archivo_activo()
    existe = activo.exists()
    info = {"existe": existe, "origen": "copia de trabajo" if RUTA_SALIDA.exists() else "planilla original"}
    if existe:
        info["modificado"] = datetime.fromtimestamp(activo.stat().st_mtime).strftime("%d/%m/%Y %H:%M")
        wb = load_workbook(activo, read_only=True)
        info["hojas"] = len([n for n in wb.sheetnames if n not in ("INDICE_CERDAS", "PLANTILLA")])
        wb.close()
    return info


def reiniciar():
    """Descarta la copia de trabajo: se vuelve a usar exel/planilla.xlsx."""
    RUTA_SALIDA.unlink(missing_ok=True)


def buscar_cerda(nro):
    """Devuelve los registros de una cerda leyendo su hoja del Excel activo."""
    archivo = archivo_activo()
    if not archivo.exists():
        raise FileNotFoundError("No hay ningún Excel cargado")
    wb = load_workbook(archivo, read_only=True)
    try:
        resultado = {"cerda": nro, "hoja": False, "en_indice": False,
                     "columnas": [], "registros": [], "indice": None}
        if "INDICE_CERDAS" in wb.sheetnames:
            for fila in wb["INDICE_CERDAS"].iter_rows(min_row=PRIMERA_FILA_INDICE, values_only=True):
                if fila and fila[0] == nro:
                    resultado["en_indice"] = True
                    resultado["indice"] = {"registros": fila[2] if len(fila) > 2 else None,
                                           "ultimo_servicio": _texto_fecha(fila[3]) if len(fila) > 3 else None}
                    break
        if str(nro) in wb.sheetnames:
            resultado["hoja"] = True
            filas = list(wb[str(nro)].iter_rows(min_row=4, values_only=True))
            if filas:
                ancho = max((i + 1 for i, c in enumerate(filas[0]) if c not in (None, "")), default=0)
                resultado["columnas"] = [str(c) for c in filas[0][:ancho]]
                for f in filas[1:]:
                    f = list(f[:ancho])
                    if any(v not in (None, "") for v in f):
                        resultado["registros"].append([_texto_fecha(v) for v in f])
        return resultado
    finally:
        wb.close()